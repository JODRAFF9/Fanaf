"""Fusionne plusieurs editions de l'annuaire FANAF dans un classeur unique.

Chaque edition couvre deux exercices (30e : 2020-2021 ; 32e : 2022-2023 ;
33e : 2023-2024...). Les tables Emission&Prestations et Chiffres cles ont une
ligne par societe, rubrique et annee, et une colonne de valeurs par edition :
une annee publiee dans deux editions a donc ses deux valeurs cote a cote.

Les societes sont rapprochees d'une edition a l'autre (meme pays et meme
branche) par le nom, la date de creation, le directeur general et le capital,
ce qui couvre les changements de nom (Atlantique -> AFG) et les noms
illisibles du PDF 2026.

Usage (sources de la plus ancienne a la plus recente) ; une source est un PDF
d'annuaire ou une liste de dossiers de formulaires Excel separes par ";" :
  python fusionner_annuaires.py sortie.xlsx "Formulaires 2020=Stes vie 2020;Stes non vie 2020" \
      "30e edition=A.pdf" "32e edition=B.pdf" "33e edition=C.pdf"
"""
import difflib
import os
from collections import Counter
import re
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import extraire_annuaire as ex
import lire_formulaires_xls as lx


def ref(f):
    """Reference de la fiche dans sa source : page du PDF ou classeur/feuille."""
    return f.get("source") or f"PDF p. {f['page']}"

PAYS_ALIAS = {"CONGO BRAZZAVILLE": "CONGO"}


# Noms de societes illisibles dans le PDF 2026 (police mal encodee), decodes a la main :
# chaque page remplace les lettres par d'autres caracteres ; la correspondance se
# retrouve a partir des mots connus ("SOCIETE", "D'ASSURANCES", "REASSURANCES"...).
NOMS_DECODES = {
    '!"#A%C%DEF*A#EA+%DI-E!!.*E+#%!D%CDI%D*%E!!.*E+#%':
        "SOCIETE AFRICAINE D'ASSURANCES ET DE REASSURANCE",
    '!"#A%CC%DE"F*%+%D*!+%I-.I*F-*!%DL0-EE#I-!F%EDDA*%':
        "NOUVELLE SOCIETE INTERAFRICAINE D'ASSURANCES VIE",
    '!"#A%C%DE)CA"E)*%D+I)!!-.)E#%!D%CD+%D.%)!!-.)E#%!DLA%':
        "SOCIETE NATIONALE D'ASSURANCES ET DE REASSURANCES VIE",
    '!"#A%C%DE)CA"E)*%D+I)!!-.)E#%!D%CD+%D.%)!!-.)E#%!':
        "SOCIETE NATIONALE D'ASSURANCES ET DE REASSURANCES",
    '!"#ABCDE)G+IB--./BD!)G)MG/)B--./BD!)G-BN.GDO.#BD':
        "COMPAGNIE D'ASSURANCE ET REASSURANCE SABU NYUMAN",
    '!"#$%&%G($)%*$%((%G+,-!!.*-(#%!G%&G+%G*%-!!.*-(#%!':
        "SOCIETE NIGERIENNE D'ASSURANCES ET DE REASSURANCES",
}


def nom_illisible(nom):
    return not nom or bool(re.search(r'[!"#$%&*+]{2}|[!"#][A-Z%]', nom))


MOTS_GENERIQUES = {"assurance", "assurances", "sa", "s", "a", "de", "des", "du", "d", "l", "la", "le", "les", "et",
                   "vie", "iard", "iardt", "cie", "compagnie", "ex", "insurance", "cote", "ivoire", "benin",
                   "burkina", "faso", "cameroun", "cameroon", "congo", "brazzaville", "gabon", "mali", "niger",
                   "senegal", "tchad", "togo", "centrafrique"}


def distinctif(nom):
    """Nom sans les mots generiques ni le pays ("SAAR ASSURANCE COTE D'IVOIRE" -> "saar")."""
    return " ".join(m for m in ex.cle(nom).split() if m not in MOTS_GENERIQUES) or ex.cle(nom)


def sim(a, b):
    return difflib.SequenceMatcher(None, distinctif(a), distinctif(b)).ratio()


def normaliser(fiches):
    for f in fiches:
        f["ident"]["Pays"] = PAYS_ALIAS.get(f["ident"]["Pays"], f["ident"]["Pays"])
        brut = f["ident"]["Societe"]
        f["decode"] = brut in NOMS_DECODES
        if f["decode"]:
            f["ident"]["Societe"] = NOMS_DECODES[brut]
        f["illisible"] = nom_illisible(f["ident"]["Societe"])
    return fiches


def rapprocher(anciennes, recentes):
    """Couples (ancienne, recente) : un a un, meilleur score d'abord."""
    candidats = []
    for i, a in enumerate(anciennes):
        for j, r in enumerate(recentes):
            if (a["ident"]["Pays"], a["branche"]) != (r["ident"]["Pays"], r["branche"]):
                continue
            nom = 0 if r["illisible"] else sim(a["ident"]["Societe"], r["ident"]["Societe"])
            date = bool(a["ident"]["Date"]) and a["ident"]["Date"] == r["ident"]["Date"]
            dg = bool(a["ident"]["DG"]) and sim(a["ident"]["DG"], r["ident"]["DG"]) > 0.8
            cap = bool(a["ident"]["Capital"]) and ex.capital(a["ident"]["Capital"]) == ex.capital(r["ident"]["Capital"])
            dates_diff = bool(a["ident"]["Date"]) and bool(r["ident"]["Date"]) and not date
            # nom quasi identique (identique si les dates de creation different),
            # ou meme date de creation confirmee par un autre indice
            if nom >= (0.99 if dates_diff else 0.85) or (date and (nom >= 0.35 or dg or cap)):
                candidats.append((nom + 0.5 * date + 0.3 * dg + 0.2 * cap, i, j))
    candidats.sort(reverse=True)
    pris_a, pris_r, couples = set(), set(), {}
    for s, i, j in candidats:
        if i not in pris_a and j not in pris_r:
            pris_a.add(i)
            pris_r.add(j)
            couples[j] = (i, s)
    return couples


def construire(editions):
    """editions : [(libelle, fiches)] de la plus ancienne a la plus recente.
    Renvoie les societes : {'fiches': {libelle: fiche}, 'nom', 'pays', 'branche'}."""
    societes, notes = [], []
    for lib, fiches in reversed(editions):
        # representant de chaque societe : sa fiche la plus recente
        reps = [s["fiches"][next(iter(s["fiches"]))] for s in societes]
        couples = rapprocher(fiches, reps)  # {indice societe: (indice fiche, score)}
        prises = {}
        for j, (i, _) in couples.items():
            prises[i] = j
        for i, f in enumerate(fiches):
            nom = f["ident"]["Societe"]
            if f.get("decode"):
                notes.append(("Nom illisible", f"{lib} : {ref(f)}", f"nom decode : {nom}"))
            if i in prises:
                soc = societes[prises[i]]
                if ex.cle(soc["nom"]) != ex.cle(nom) and not f["illisible"]:
                    notes.append(("Changement de nom", f"{lib} : {ref(f)}", f"{nom} -> {soc['nom']}"))
                soc["fiches"][lib] = f
            else:
                if f["illisible"]:
                    nom = f"(nom illisible, {lib} : {ref(f)})"
                    notes.append(("Nom illisible", f"{lib} : {ref(f)}", "aucune correspondance trouvee"))
                societes.append({"fiches": {lib: f}, "nom": nom, "pays": f["ident"]["Pays"], "branche": f["branche"]})
    # nom illisible dans l'edition recente mais lisible dans une plus ancienne
    for soc in societes:
        if soc["nom"].startswith("(nom illisible"):
            lisibles = [f for f in soc["fiches"].values() if not f["illisible"]]
            if lisibles:
                notes.append(("Nom illisible", soc["nom"], f"nom repris d'une edition anterieure : {lisibles[0]['ident']['Societe']}"))
                soc["nom"] = lisibles[0]["ident"]["Societe"]
    societes.sort(key=lambda s: (s["pays"], s["branche"] != "Vie", ex.cle(s["nom"])))
    return societes, notes


def ecrire(societes, notes, sortie, libelles):
    wb = Workbook()
    wb.remove(wb.active)
    gras = Font(bold=True)

    def feuille(nom, entetes, cache=False):
        ws = wb.create_sheet(nom)
        ws.append(entetes)
        for c in ws[1]:
            c.font = gras
        ws.freeze_panes = "A2"
        if cache:
            ws.sheet_state = "hidden"
        return ws

    e_id = ["No enregistrement", "No societe", "Edition", "Societe", "Pays", "Branche", "Directeur general",
            "Date de creation", "Capital social (F CFA)", "Cadres", "Maitrise", "Employes", "Annee N", "Source"]
    idv = feuille("Identification", e_id)
    e_ep = ["No societe", "Societe", "Pays", "Branche", "Bloc", "Categorie", "Rubrique", "Mesure", "Annee"]
    e_cc = ["No societe", "Societe", "Pays", "Branche", "Rubrique", "Annee"]
    ne = len(libelles)
    epv = feuille("Emission&Prestations", e_ep + [f"Valeur {e} (F CFA)" for e in libelles])
    ccv = feuille("Chiffres clés", e_cc + [f"Valeur {e} (F CFA)" for e in libelles])
    epb = feuille("Emission&Prestations (brut)", e_ep + [f"Valeur {e} (milliers F CFA)" for e in libelles]
                  + [f"Source {e}" for e in libelles], True)
    ccb = feuille("Chiffres clés (brut)", e_cc + [f"Valeur {e} (milliers F CFA)" for e in libelles]
                  + [f"Source {e}" for e in libelles], True)
    col_ep = [get_column_letter(len(e_ep) + 1 + k) for k in range(ne)]
    col_cc = [get_column_letter(len(e_cc) + 1 + k) for k in range(ne)]

    n_enr = 0
    for no, s in enumerate(societes, 1):
        base_soc = [no, s["nom"], s["pays"], s["branche"]]
        # Identification : une ligne par fiche (edition)
        for lib in libelles:
            f = s["fiches"].get(lib)
            if f is None:
                continue
            n_enr += 1
            i = f["ident"]
            idv.append([n_enr, no, lib, s["nom"] if f["illisible"] else i["Societe"], s["pays"],
                        s["branche"], i["DG"] or None, ex.date_ou_texte(i["Date"]), ex.capital(i["Capital"]),
                        i["Cadres"], i["Maitrise"], i["Employes"], f["annees"][1], ref(f)])
        # Valeurs : cle -> [valeur par edition] + [source par edition]
        lignes_ep, lignes_cc = {}, {}
        for k, lib in enumerate(libelles):
            f = s["fiches"].get(lib)
            if f is None:
                continue
            for typ, bloc, cat, rub, mes, an, v in f["entrees"]:
                if typ == "EP":
                    d, cle_l = lignes_ep, (bloc, cat or None, rub, mes, an)
                else:
                    d, cle_l = lignes_cc, (rub, an)
                e = d.setdefault(cle_l, [None] * (2 * ne))
                e[k] = v
                e[ne + k] = ref(f)
        for d, wsb, wsv, cols, nomb in ((lignes_ep, epb, epv, col_ep, "Emission&Prestations (brut)"),
                                        (lignes_cc, ccb, ccv, col_cc, "Chiffres clés (brut)")):
            for cle_l in sorted(d, key=lambda c: (c[:-1], c[-1])):
                wsb.append(base_soc + list(cle_l) + d[cle_l])
                r = wsb.max_row
                wsv.append(base_soc + list(cle_l) + [
                    f"=IF('{nomb}'!{c}{r}=\"\",\"\",'{nomb}'!{c}{r}*1000)" for c in cols])

    nt = feuille("Notes", ["Type", "Reference", "Detail"])
    for n in notes:
        nt.append(list(n))

    for c in idv["H"][1:]:
        c.number_format = "dd/mm/yyyy"
    for c in idv["I"][1:]:
        c.number_format = "#,##0"
    for ws, cols in ((epv, col_ep), (ccv, col_cc)):
        for col in cols:
            for c in ws[col][1:]:
                c.number_format = "#,##0"
    for ws, cols in ((epb, col_ep), (ccb, col_cc)):
        for col in cols:
            for c in ws[col][1:]:
                c.number_format = "#,##0.000"
    for ws in wb.worksheets:
        for col in ws.columns:
            larg = max((len(str(c.value)) for c in col[:300]
                        if c.value is not None and not str(c.value).startswith("=")), default=12)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, larg + 2), 60)
    wb.save(sortie)


def main(sortie, editions_pdf):
    """editions_pdf : [(libelle, chemin du PDF)] de la plus ancienne a la plus recente."""
    editions = []
    notes_lecture = []
    for lib, chemin in editions_pdf:
        if ";" in chemin or os.path.isdir(chemin):
            fs, ign = lx.lire_dossiers([c for c in chemin.split(";") if c])  # formulaires Excel
        else:
            fs, ign = ex.lire_pdf(chemin)
        normaliser(fs)
        editions.append((lib, fs))
        if ign:
            notes_lecture.append(("Non lues", lib, ", ".join(map(str, ign))))
        attendu = Counter(f["annees"] for f in fs).most_common(1)[0][0] if fs else None
        for f in fs:
            for c in f["controles"]:
                notes_lecture.append(("Controle", f"{lib} : {ref(f)} {f['ident']['Societe'][:60]}", c))
            if f["annees"] != attendu:
                notes_lecture.append(("Exercices", f"{lib} : {ref(f)} {f['ident']['Societe']}",
                                      f"fiche publiee avec les exercices {f['annees'][0]}-{f['annees'][1]}"))
    societes, notes = construire(editions)
    notes += notes_lecture
    libelles = [lib for lib, _ in editions_pdf]
    ecrire(societes, notes, sortie, libelles)
    for lib, fs in editions:
        print(f"{lib} : {len(fs)} fiches")
    print(f"{len(societes)} societes")
    return societes, notes


if __name__ == "__main__":
    main(sys.argv[1], [tuple(a.split("=", 1)) for a in sys.argv[2:]])
