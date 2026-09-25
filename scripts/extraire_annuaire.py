"""Extrait les donnees brutes des fiches societes de l'annuaire FANAF (PDF)
vers un classeur Excel ayant la structure des tables du module VBA
ModConsolidation : Identification, Emission&Prestations, Chiffres cles
(feuilles visibles + feuilles "(brut)" masquees).

Seules les donnees saisies sont reprises ; les valeurs deduites par formule
(totaux, Ensemble, parts, evolutions, ratios CS/PA, autres actifs, taux de
couverture, nombre total de salaries) ne le sont pas. Elles servent a
controler la lecture : la somme des branches doit redonner le total du PDF.

Usage : python extraire_annuaire.py ANNUAIRE.pdf sortie.xlsx
"""
import difflib
import re
import sys
import unicodedata
from datetime import datetime

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName

ANNEE = re.compile(r"^(19|20)\d\d$")
POURCENT = re.compile(r"%\s*$")
FRAGMENT = re.compile(r"^-?[\d ]+(,\d+)?$|^-$")

# Libelles de reference (formulaires FANAF)
RUBRIQUES_VIE = ["Contrats en cas de vie", "Contrat en cas de décès", "Mixte", "Epargne",
                 "Titre de capitalisation", "Complémentaires", "Acceptations"]
BRANCHES_NV = ["Accidents corporels et maladie", "Automobile : Responsabilité Civile",
               "Automobile : Autres risques", "Incendie et autres dommages aux biens", "RC Générale",
               "Transports Aériens", "Transports Maritimes", "Autres Transports", "Autres Risques",
               "Acceptations"]
RUBRIQUES_C = ["Produits financiers nets", "Commissions", "Autres charges de l'exercice",
               "Provisions mathématiques (clôture)", "Provisions mathématiques (ouverture)",
               "Provisions pour risques en cours", "Primes cédées aux réassureurs",
               "Primes acquises aux réassureurs", "Part des réassureurs dans les charges",
               "Prestations et frais payés", "Provisions pour sinistres à payer",
               "Résultats d’exploitations nets", "Total des capitaux propres et réserves",
               "Résultats au Bilan", "Marge réglementaire", "Marge disponible", "Engagements Réglementés",
               "Actifs admis", "Dont liquidités", "Autres actifs"]


# Libelles tronques ou alteres dans le PDF (p. 86 : "Actifs" ; p. 107 : "U°E")
ALIAS = {"actifs": "Actifs admis", "ue": "Accidents corporels et maladie",
         # anciens libelles, verifies sur les valeurs 2020 (formulaires = 30e edition)
         "frais generaux": "Autres charges de l'exercice",
         "fonds propres nets": "Total des capitaux propres et réserves"}


def cle(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    s = s.replace("’", "'")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def canonique(lab, reference, seuil=0.84):
    """Rapproche un libelle de la liste de reference (fautes, pluriels, accents)."""
    k = cle(lab)
    if k in ALIAS and ALIAS[k] in reference:
        return ALIAS[k]
    cles = {cle(r): r for r in reference}
    if k in cles:
        return cles[k]
    m = difflib.get_close_matches(k, cles.keys(), n=1, cutoff=seuil)
    return cles[m[0]] if m else lab


def valeur(t):
    """Nombre lu dans un texte du PDF ; "-" vaut 0 ; None si ce n'est pas un nombre."""
    t = t.strip()
    if t == "-":
        return 0.0
    neg = t.startswith("-")
    t = t.lstrip("-").strip().replace(" ", "")
    if not re.fullmatch(r"\d+(,\d+)?", t):
        return None
    v = float(t.replace(",", "."))
    return -v if neg else v


def est_nombre(t):
    return not POURCENT.search(t) and valeur(t) is not None


def libelle(t):
    s = " ".join(t.split())
    return s[:1].upper() + s[1:]


def lignes(page):
    """Mots regroupes en lignes ; positions calculees sur les caracteres non blancs ;
    fragments de nombre contigus fusionnes ("9" + "5 037" -> "95 037")."""
    mots = []
    for w in page.extract_words(keep_blank_chars=True, x_tolerance=1.5, y_tolerance=2, return_chars=True):
        cs = [c for c in w["chars"] if c["text"].strip()]
        if not cs:
            continue
        mots.append({"text": w["text"].strip(), "x0": min(c["x0"] for c in cs),
                     "x1": max(c["x1"] for c in cs), "top": w["top"]})
    mots.sort(key=lambda m: (m["top"], m["x0"]))
    res, cur, y = [], [], None
    for m in mots:
        if y is None or abs(m["top"] - y) > 2.5:
            if cur:
                res.append(cur)
            cur, y = [], m["top"]
        cur.append(m)
    if cur:
        res.append(cur)
    fusion = []
    for l in res:
        l = sorted(l, key=lambda m: m["x0"])
        out = []
        for m in l:
            if (out and FRAGMENT.match(out[-1]["text"]) and FRAGMENT.match(m["text"])
                    and m["x0"] - out[-1]["x1"] < 4 and out[-1]["text"] != "-" or
                    out and out[-1]["text"] == "-" and FRAGMENT.match(m["text"]) and m["x0"] - out[-1]["x1"] < 60
                    and out[-1]["x0"] > 150):
                out[-1] = {"text": out[-1]["text"] + " " + m["text"], "x0": out[-1]["x0"],
                           "x1": m["x1"], "top": out[-1]["top"]}
            else:
                out.append(dict(m))
        fusion.append(out)
    return fusion


def centre(m):
    return (m["x0"] + m["x1"]) / 2


def decouper(ligne, x_label):
    lab = " ".join(m["text"] for m in ligne if m["x0"] < x_label and not est_nombre(m["text"])
                   and not POURCENT.search(m["text"]))
    montants = [m for m in ligne if m["x0"] >= x_label - 60 and est_nombre(m["text"])]
    return libelle(lab), montants


def affecter(montants, centres):
    res = [None] * len(centres)
    for m in montants:
        i = min(range(len(centres)), key=lambda k: abs(centre(m) - centres[k]))
        res[i] = valeur(m["text"])
    return res


def texte_apres(ligne, x_min=150):
    return " ".join(m["text"] for m in ligne if m["x0"] >= x_min).strip()


def lire_fiche(page, num_page):
    ls = lignes(page)
    txt = [" ".join(m["text"] for m in l) for l in ls]

    def index(motif, debut=0, fin=None):
        for i in range(debut, fin if fin is not None else len(ls)):
            if re.search(motif, txt[i], re.I):
                return i
        return None

    iA = index(r"(^| )[AB]? ?- ?EMISSIONS NETTES")
    iB = index(r"(^| )[ABC]? ?- ?(PRESTATIONS|SINISTRALITE)", iA + 1 if iA is not None else 0)
    if iA is None or iB is None:
        return None
    iC = index(r"^Rubriques", iB)
    if iC is None:
        return None
    branche = "Non-vie" if re.search("SINISTRALITE", txt[iB], re.I) else "Vie"

    # --- Identification ---
    ident = {"Societe": "", "Pays": "", "DG": "", "Date": "", "Capital": "",
             "Cadres": None, "Maitrise": None, "Employes": None}
    noms_eff = {"cadres": "Cadres", "maîtrise": "Maitrise", "maitrise": "Maitrise",
                "employés": "Employes", "employes": "Employes"}
    for l, t in zip(ls[:iA], txt[:iA]):
        tl = t.lower()
        if tl.startswith("nom de la soci"):
            ident["Societe"] = texte_apres(l)
        elif tl.startswith("pays"):
            ident["Pays"] = texte_apres(l)
        elif tl.startswith("nom du directeur") or tl.startswith("nom ad"):
            ident["DG"] = texte_apres(l)
        elif tl.startswith("date de cr"):
            ident["Date"] = texte_apres(l)
        elif tl.startswith("capital social") or tl.startswith("fonds d"):
            ident["Capital"] = texte_apres(l)
        for m in l:
            k = m["text"].lower()
            if k in noms_eff:
                vals = [valeur(x["text"]) for x in l if x["x0"] > m["x1"] and est_nombre(x["text"])]
                ident[noms_eff[k]] = vals[0] if vals else None

    # Annees : 1re ligne apres le titre A portant deux annees
    annees, ent_A = None, None
    for l in ls[iA:iB]:
        ys = [m for m in l if ANNEE.match(m["text"])]
        if len(ys) >= 2:
            annees, ent_A = (int(ys[0]["text"]), int(ys[1]["text"])), ys[:2]
            break
    if not annees:
        return None

    entrees, controles = [], []

    def bloc_ab(i0, i1, nom_bloc):
        blk = ls[i0 + 1:i1]
        # colonnes des montants
        ent = next((l for l in blk if sum(1 for m in l if ANNEE.match(m["text"])) >= 2), None)
        cent_annees = [centre(m) for m in ent if ANNEE.match(m["text"])][:2] if ent else [centre(m) for m in ent_A]
        if branche == "Non-vie" and nom_bloc == "Emissions":
            mesures = ["Primes émises", "Primes acquises (PA)"]
            ref = next((decouper(l, 250)[1] for l in blk
                        if re.match(r"(ensemble|total|chiffre d)", " ".join(m["text"] for m in l), re.I)
                        and len(decouper(l, 250)[1]) == 4), None)
            if ref:
                centres = [centre(m) for m in ref]
            else:
                sous = [m for l in blk[:4] for m in l if re.match(r"primes\b", m["text"], re.I)]
                sous.sort(key=lambda m: m["x0"])
                if len(sous) < 4:
                    controles.append(f"{nom_bloc} : colonnes introuvables")
                    return
                centres = [centre(m) for m in sous[:4]]
            cols = [(mesures[0], annees[0]), (mesures[1], annees[0]), (mesures[0], annees[1]), (mesures[1], annees[1])]
        else:
            mes = ("Charges de sinistres (CS)" if branche == "Non-vie" else
                   "Emissions nettes" if nom_bloc == "Emissions" else "Prestations versees")
            centres = cent_annees
            cols = [(mes, annees[0]), (mes, annees[1])]

        reference = BRANCHES_NV if branche == "Non-vie" else RUBRIQUES_VIE
        categorie = "Affaires directes" if branche == "Non-vie" else ""
        n_vie = 0
        sommes, totaux = {}, {}
        for l in blk:
            lab, mont = decouper(l, 250 if branche == "Non-vie" else 215)
            seul = " ".join(m["text"] for m in l)
            if re.match(r"assurances (individuelles|collectives)", seul, re.I):
                categorie = libelle(seul.split(None, 1)[1])
                continue
            low = lab.lower()
            if (not lab or low.startswith("branches") or low.startswith("(chiffre")
                    or re.match(r"(\(cs\)|primes)", low)):
                continue
            vals = affecter(mont, centres)
            if low.startswith("total") or low.startswith("ensemble") or low.startswith("chiffre d"):
                cle_t = "total" if low.startswith("total") else "ensemble"
                for (mes, an), v in zip(cols, vals):
                    totaux[(cle_t, mes, an)] = v
                continue
            lab = canonique(lab, reference)
            if branche == "Vie" and cle(lab).startswith("contrat") and cle(lab).endswith("de vie"):
                n_vie += 1
                categorie = "Individuelles" if n_vie == 1 else "Collectives"
            if lab == "Contrats en cas de vie" and categorie == "Collectives":
                lab = "Contrat en cas de vie"
            cat = "Acceptations" if lab == "Acceptations" else categorie
            for (mes, an), v in zip(cols, vals):
                entrees.append(("EP", nom_bloc, cat, lab, mes, an, v))
                k = ("acc", mes, an) if cat == "Acceptations" else (mes, an)
                sommes[k] = sommes.get(k, 0) + (v or 0)
        for (mes, an) in cols:
            s = sommes.get((mes, an), 0)
            t = totaux.get(("total", mes, an))
            e = totaux.get(("ensemble", mes, an))
            acc = sommes.get(("acc", mes, an), 0)
            if t is not None and abs(s - t) > 2:
                controles.append(f"{nom_bloc} / {mes} / {an} : somme des branches {s:,.0f} "
                                 f"<> total affaires directes {t:,.0f}".replace(",", " "))
            if e is not None and abs(s + acc - e) > 2:
                controles.append(f"{nom_bloc} / {mes} / {an} : branches + acceptations {s + acc:,.0f} "
                                 f"<> ensemble {e:,.0f}".replace(",", " "))

    bloc_ab(iA, iB, "Emissions")
    bloc_ab(iB, iC, "Sinistralite" if branche == "Non-vie" else "Prestations")

    # --- Bloc C ---
    ent = ls[iC] if sum(1 for m in ls[iC] if ANNEE.match(m["text"])) >= 2 else \
        next((l for l in ls[iC:iC + 2] if sum(1 for m in l if ANNEE.match(m["text"])) >= 2), None)
    centres = [centre(m) for m in ent if ANNEE.match(m["text"])][:2] if ent else None
    if centres is None:
        # en-tete incomplet (une seule annee) : colonnes deduites de la position des montants,
        # separees au plus grand ecart entre les centres
        xs = sorted(centre(m) for l in ls[iC + 1:iC + 12] for m in decouper(l, 260)[1])
        if len(xs) >= 4:
            k = max(range(1, len(xs)), key=lambda q: xs[q] - xs[q - 1])
            centres = [sum(xs[:k]) / k, sum(xs[k:]) / (len(xs) - k)]
            controles.append("Chiffres cles : en-tete incomplet, colonnes deduites des montants")
    if centres is None:
        controles.append("Chiffres cles : en-tete introuvable")
    else:
        vals_c = {}
        for l in ls[iC + 1:]:
            lab, mont = decouper(l, 260)
            if not lab or re.fullmatch(r"-?\s*\d+\s*-?", lab):
                continue
            if lab.lower().startswith("taux"):
                break
            lab = canonique(lab, RUBRIQUES_C)
            v = affecter(mont, centres)
            vals_c[lab] = v
            if lab == "Autres actifs":
                continue
            for k, an in enumerate(annees):
                entrees.append(("CC", "", "", lab, "", an, v[k]))
        aa, liq, au = vals_c.get("Actifs admis"), vals_c.get("Dont liquidités"), vals_c.get("Autres actifs")
        if aa and liq and au:
            for k, an in enumerate(annees):
                if None not in (aa[k], liq[k], au[k]) and abs(aa[k] - liq[k] - au[k]) > 2:
                    controles.append(f"Chiffres cles / {an} : actifs admis - liquidites "
                                     f"{aa[k] - liq[k]:,.0f} <> autres actifs {au[k]:,.0f}".replace(",", " "))

    return {"page": num_page, "page_imprimee": next((t.strip(" -") for t in txt[::-1]
                                                     if re.fullmatch(r"-\s*\d+\s*-", t.strip())), ""),
            "branche": branche, "annees": annees, "ident": ident, "entrees": entrees, "controles": controles}


def date_ou_texte(s):
    try:
        return datetime.strptime(s.strip(), "%d/%m/%Y")
    except ValueError:
        return s or None


def capital(s):
    ch = re.sub(r"\D", "", s or "")
    return float(ch) if ch else (s or None)


def ecrire(fiches, sortie):
    wb = Workbook()
    wb.remove(wb.active)
    gras = Font(bold=True)

    def feuille(nom, entetes, cache=False):
        ws = wb.create_sheet(nom)
        ws.append(entetes)
        for c in ws[1]:
            c.font = gras
        ws.freeze_panes = "A2"
        if cache:
            ws.sheet_state = "hidden"
        return ws

    e_id = ["No enregistrement", "Societe", "Pays", "Branche", "Directeur general", "Date de creation",
            "Capital social (F CFA)", "Cadres", "Maitrise", "Employes", "Annee N"]
    e_ep = ["No enregistrement", "Societe", "Pays", "Branche", "Bloc", "Categorie", "Rubrique", "Mesure", "Annee"]
    e_cc = ["No enregistrement", "Societe", "Pays", "Branche", "Rubrique", "Annee"]
    idv = feuille("Identification", e_id)
    epv = feuille("Emission&Prestations", e_ep + ["Valeur (F CFA)"])
    ccv = feuille("Chiffres clés", e_cc + ["Valeur (F CFA)"])
    idb = feuille("Identification (brut)", e_id + ["Date d'import"], True)
    epb = feuille("Emission&Prestations (brut)", e_ep + ["Valeur (milliers F CFA)", "Cellule source"], True)
    ccb = feuille("Chiffres clés (brut)", e_cc + ["Valeur (milliers F CFA)", "Cellule source"], True)

    maintenant = datetime.now().replace(microsecond=0)
    for n, f in enumerate(fiches, 1):
        i = f["ident"]
        ligne_id = [n, i["Societe"], i["Pays"], f["branche"], i["DG"] or None, date_ou_texte(i["Date"]),
                    capital(i["Capital"]), i["Cadres"], i["Maitrise"], i["Employes"], f["annees"][1]]
        idb.append(ligne_id + [maintenant])
        idv.append(ligne_id)
        src = f"PDF p. {f['page']}" + (f" (page imprimee {f['page_imprimee']})" if f["page_imprimee"] else "")
        for typ, bloc, cat, rub, mes, an, v in f["entrees"]:
            if typ == "EP":
                base = [n, i["Societe"], i["Pays"], f["branche"], bloc, cat or None, rub, mes, an]
                epb.append(base + [v, src])
                r = epb.max_row
                epv.append(base + [f"=IF('Emission&Prestations (brut)'!J{r}=\"\",\"\","
                                   f"'Emission&Prestations (brut)'!J{r}*1000)"])
            else:
                base = [n, i["Societe"], i["Pays"], f["branche"], rub, an]
                ccb.append(base + [v, src])
                r = ccb.max_row
                ccv.append(base + [f"=IF('Chiffres clés (brut)'!G{r}=\"\",\"\","
                                   f"'Chiffres clés (brut)'!G{r}*1000)"])

    for ws in (idv, idb):
        for c in ws["F"][1:]:
            c.number_format = "dd/mm/yyyy"
        for c in ws["G"][1:]:
            c.number_format = "#,##0"
    for c in idb["L"][1:]:
        c.number_format = "dd/mm/yyyy hh:mm"
    for ws, col in ((epv, "J"), (ccv, "G")):
        for c in ws[col][1:]:
            c.number_format = "#,##0"
    for ws, col in ((epb, "J"), (ccb, "G")):
        for c in ws[col][1:]:
            c.number_format = "#,##0.000"
    for ws in wb.worksheets:
        for col in ws.columns:
            larg = max((len(str(c.value)) for c in col[:300]
                        if c.value is not None and not str(c.value).startswith("=")), default=12)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, larg + 2), 55)

    # Compteur du module VBA : les prochains enregistrements continuent la numerotation
    wb.defined_names["NumeroEnregistrement"] = DefinedName("NumeroEnregistrement",
                                                           attr_text=str(len(fiches)), hidden=True)
    wb.save(sortie)


def pays_normalise(p):
    return unicodedata.normalize("NFKD", p).encode("ascii", "ignore").decode().upper().strip()


def lire_pdf(pdf_path):
    """Lit toutes les fiches societes du PDF. Renvoie (fiches, pages non lues)."""
    fiches, ignorees = [], []
    pays_section = ""
    with pdfplumber.open(pdf_path) as pdf:
        for k, page in enumerate(pdf.pages, 1):
            t = page.extract_text() or ""
            # page de titre d'un pays : son seul texte est le nom du pays
            mots = [l.strip() for l in t.splitlines() if l.strip() and not re.fullmatch(r"-\s*\d+\s*-", l.strip())]
            if k > 10 and len(mots) == 1 and mots[0].isupper() and len(mots[0]) < 30:
                pays_section = pays_normalise(mots[0])
                continue
            if not re.search(r"NOM DE LA SOCIETE", t) or not re.search(r"EMISSIONS NETTES", t):
                continue
            f = lire_fiche(page, k)
            if f is None:
                ignorees.append(k)
                continue
            p = pays_normalise(f["ident"]["Pays"])
            if not re.search(r"[A-Z]", p) and pays_section:
                f["controles"].append(f"pays lu dans le PDF : '{f['ident']['Pays']}', remplace par {pays_section}")
                p = pays_section
            f["ident"]["Pays"] = p
            fiches.append(f)
    return fiches, ignorees


def main(pdf_path, sortie):
    fiches, ignorees = lire_pdf(pdf_path)
    ecrire(fiches, sortie)
    print(f"{len(fiches)} fiches extraites ({sum(f['branche'] == 'Vie' for f in fiches)} vie, "
          f"{sum(f['branche'] == 'Non-vie' for f in fiches)} non-vie)")
    if ignorees:
        print("Pages non lues :", ignorees)
    for f in fiches:
        manques = [k for k in ("Societe", "Pays", "DG", "Date", "Capital") if not f["ident"][k]]
        if f["controles"] or manques:
            print(f"p. {f['page']} {f['ident']['Societe'] or '(sans nom)'} :")
            if manques:
                print("   champs vides :", ", ".join(manques))
            for c in f["controles"]:
                print("   " + c)
    return fiches


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
