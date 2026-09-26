"""Fusionne plusieurs editions de l'annuaire FANAF dans un classeur unique.

Chaque edition couvre deux exercices (30e : 2020-2021 ; 32e : 2022-2023 ;
33e : 2023-2024...). Les tables Emission&Prestations et Chiffres cles ont une
ligne par societe, rubrique et annee, et une colonne de valeurs par edition :
une annee publiee dans deux editions a donc ses deux valeurs cote a cote.

Les societes sont rapprochees d'une edition a l'autre (meme pays et meme
branche) par le nom, la date de creation, le directeur general et le capital,
ce qui couvre les changements de nom (Atlantique -> AFG) et les noms
illisibles du PDF 2026.

Usage (sources de la plus ancienne a la plus recente) ; une source est un PDF
d'annuaire ou une liste de dossiers de formulaires Excel separes par ";" :
  python fusionner_annuaires.py sortie.xlsx "Formulaires 2020=Stes vie 2020;Stes non vie 2020" \
      "30e edition=A.pdf" "32e edition=B.pdf" "33e edition=C.pdf"
"""
import difflib
import os
from collections import Counter
import re
import sys
from datetime import datetime
from multiprocessing import Pool

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import extraire_annuaire as ex
import lire_formulaires_xls as lx
import verifier_extraction as vx


def ref(f):
    """Reference de la fiche dans sa source : page du PDF ou classeur/feuille."""
    return f.get("source") or f"PDF p. {f['page']}"

PAYS_ALIAS = {"CONGO BRAZZAVILLE": "CONGO"}


# Noms de societes illisibles dans le PDF 2026 (police mal encodee), decodes a la main :
# chaque page remplace les lettres par d'autres caracteres ; la correspondance se
# retrouve a partir des mots connus ("SOCIETE", "D'ASSURANCES", "REASSURANCES"...).
NOMS_DECODES = {
    '!"#A%C%DEF*A#EA+%DI-E!!.*E+#%!D%CDI%D*%E!!.*E+#%':
        "SOCIETE AFRICAINE D'ASSURANCES ET DE REASSURANCE",
    '!"#A%CC%DE"F*%+%D*!+%I-.I*F-*!%DL0-EE#I-!F%EDDA*%':
        "NOUVELLE SOCIETE INTERAFRICAINE D'ASSURANCES VIE",
    '!"#A%C%DE)CA"E)*%D+I)!!-.)E#%!D%CD+%D.%)!!-.)E#%!DLA%':
        "SOCIETE NATIONALE D'ASSURANCES ET DE REASSURANCES VIE",
    '!"#A%C%DE)CA"E)*%D+I)!!-.)E#%!D%CD+%D.%)!!-.)E#%!':
        "SOCIETE NATIONALE D'ASSURANCES ET DE REASSURANCES",
    '!"#ABCDE)G+IB--./BD!)G)MG/)B--./BD!)G-BN.GDO.#BD':
        "COMPAGNIE D'ASSURANCE ET REASSURANCE SABU NYUMAN",
    '!"#$%&%G($)%*$%((%G+,-!!.*-(#%!G%&G+%G*%-!!.*-(#%!':
        "SOCIETE NIGERIENNE D'ASSURANCES ET DE REASSURANCES",
}


def nom_illisible(nom):
    return not nom or bool(re.search(r'[!"#$%&*+]{2}|[!"#][A-Z%]', nom))


MOTS_GENERIQUES = {"assurance", "assurances", "sa", "s", "a", "de", "des", "du", "d", "l", "la", "le", "les", "et",
                   "vie", "iard", "iardt", "cie", "compagnie", "ex", "insurance", "cote", "ivoire", "benin",
                   "burkina", "faso", "cameroun", "cameroon", "congo", "brazzaville", "gabon", "mali", "niger",
                   "senegal", "tchad", "togo", "centrafrique"}


def distinctif(nom):
    """Nom sans les mots generiques ni le pays ("SAAR ASSURANCE COTE D'IVOIRE" -> "saar")."""
    return " ".join(m for m in ex.cle(nom).split() if m not in MOTS_GENERIQUES) or ex.cle(nom)


def sim(a, b):
    return difflib.SequenceMatcher(None, distinctif(a), distinctif(b)).ratio()


def normaliser(fiches):
    for f in fiches:
        f["ident"]["Pays"] = PAYS_ALIAS.get(f["ident"]["Pays"], f["ident"]["Pays"])
        brut = f["ident"]["Societe"]
        f["decode"] = brut in NOMS_DECODES
        if f["decode"]:
            f["ident"]["Societe"] = NOMS_DECODES[brut]
        f["illisible"] = nom_illisible(f["ident"]["Societe"])
    return fiches


def rapprocher(anciennes, recentes):
    """Couples (ancienne, recente) : un a un, meilleur score d'abord."""
    candidats = []
    for i, a in enumerate(anciennes):
        for j, r in enumerate(recentes):
            if (a["ident"]["Pays"], a["branche"]) != (r["ident"]["Pays"], r["branche"]):
                continue
            nom = 0 if r["illisible"] else sim(a["ident"]["Societe"], r["ident"]["Societe"])
            date = bool(a["ident"]["Date"]) and a["ident"]["Date"] == r["ident"]["Date"]
            dg = bool(a["ident"]["DG"]) and sim(a["ident"]["DG"], r["ident"]["DG"]) > 0.8
            cap = bool(a["ident"]["Capital"]) and ex.capital(a["ident"]["Capital"]) == ex.capital(r["ident"]["Capital"])
            dates_diff = bool(a["ident"]["Date"]) and bool(r["ident"]["Date"]) and not date
            # nom quasi identique (identique si les dates de creation different),
            # ou meme date de creation confirmee par un autre indice
            if nom >= (0.99 if dates_diff else 0.85) or (date and (nom >= 0.35 or dg or cap)):
                candidats.append((nom + 0.5 * date + 0.3 * dg + 0.2 * cap, i, j))
    candidats.sort(reverse=True)
    pris_a, pris_r, couples = set(), set(), {}
    for s, i, j in candidats:
        if i not in pris_a and j not in pris_r:
            pris_a.add(i)
            pris_r.add(j)
            couples[j] = (i, s)
    return couples


def construire(editions):
    """editions : [(libelle, fiches)] de la plus ancienne a la plus recente.
    Renvoie les societes : {'fiches': {libelle: fiche}, 'nom', 'pays', 'branche'}."""
    societes, notes = [], []
    for lib, fiches in reversed(editions):
        # representant de chaque societe : sa fiche la plus recente
        reps = [s["fiches"][next(iter(s["fiches"]))] for s in societes]
        couples = rapprocher(fiches, reps)  # {indice societe: (indice fiche, score)}
        prises = {}
        for j, (i, _) in couples.items():
            prises[i] = j
        for i, f in enumerate(fiches):
            nom = f["ident"]["Societe"]
            if f.get("decode"):
                notes.append(("Nom illisible", f"{lib} : {ref(f)}", f"nom decode : {nom}"))
            if i in prises:
                soc = societes[prises[i]]
                if ex.cle(soc["nom"]) != ex.cle(nom) and not f["illisible"]:
                    notes.append(("Changement de nom", f"{lib} : {ref(f)}", f"{nom} -> {soc['nom']}"))
                soc["fiches"][lib] = f
            else:
                if f["illisible"]:
                    nom = f"(nom illisible, {lib} : {ref(f)})"
                    notes.append(("Nom illisible", f"{lib} : {ref(f)}", "aucune correspondance trouvee"))
                societes.append({"fiches": {lib: f}, "nom": nom, "pays": f["ident"]["Pays"], "branche": f["branche"]})
    # nom illisible dans l'edition recente mais lisible dans une plus ancienne
    for soc in societes:
        if soc["nom"].startswith("(nom illisible"):
            lisibles = [f for f in soc["fiches"].values() if not f["illisible"]]
            if lisibles:
                notes.append(("Nom illisible", soc["nom"], f"nom repris d'une edition anterieure : {lisibles[0]['ident']['Societe']}"))
                soc["nom"] = lisibles[0]["ident"]["Societe"]
    societes.sort(key=lambda s: (s["pays"], s["branche"] != "Vie", ex.cle(s["nom"])))
    return societes, notes


ND = "Non disponible"


def nomenclature(f):
    """Groupe de mise en page d'une fiche : branche, et pour la vie, ancienne ou nouvelle nomenclature."""
    if f["branche"] != "Vie":
        return f["branche"]
    anciennes = set(ex.RUBRIQUES_VIE_ANCIENNES) - {"Complémentaires"}
    return "Vie (ancienne nomenclature)" if any(e[0] == "EP" and not e[2] and e[3] in anciennes
                                                for e in f["entrees"]) else "Vie"


def structures(editions, seuil=0.2):
    """Rubriques publiees par chaque edition, par groupe de mise en page : une cle figure dans
    la structure si au moins 20 % des fiches du groupe la portent (case vide comprise).
    Renvoie ({(edition, groupe): cles}, {branche: toutes les cles connues})."""
    par_groupe, union = {}, {}
    for lib, fiches in editions:
        compte, n = {}, Counter()
        for f in fiches:
            g = nomenclature(f)
            n[g] += 1
            for c in {e[:5] for e in f["entrees"]}:
                compte[(g, c)] = compte.get((g, c), 0) + 1
        for (g, c), k in compte.items():
            if k >= seuil * n[g]:
                par_groupe.setdefault((lib, g), set()).add(c)
    for (lib, g), cles in par_groupe.items():
        union.setdefault("Vie" if g.startswith("Vie") else g, set()).update(cles)
    return par_groupe, union


def ecrire(societes, notes, sortie, libelles, bilans=None, editions=None):
    wb = Workbook()
    wb.remove(wb.active)
    gras = Font(bold=True)

    def feuille(nom, entetes, cache=False):
        ws = wb.create_sheet(nom)
        ws.append(entetes)
        for c in ws[1]:
            c.font = gras
        ws.freeze_panes = "A2"
        if cache:
            ws.sheet_state = "hidden"
        return ws

    par_groupe, union = structures(editions or [])
    rang = {lib: k for k, lib in enumerate(libelles)}

    # Valeurs : (type, cle) -> {annee: [(edition, valeur, reference)]}, dans l'ordre des editions.
    # Une rubrique absente de la mise en page d'une edition y vaut "Non disponible".
    donnees = []
    npub = 1
    for s in societes:
        d = {}
        for lib in libelles:
            f = s["fiches"].get(lib)
            if f is None:
                continue
            lus = {}
            for typ, bloc, cat, rub, mes, an, v in f["entrees"]:
                lus[(typ, bloc, cat, rub, mes, an)] = v
            structure = par_groupe.get((lib, nomenclature(f)), set())
            for c in union.get(f["branche"], ()):
                if c not in structure and c[2] != "Total":
                    for an in f["annees"]:
                        lus.setdefault(c + (an,), ND)
            for (typ, bloc, cat, rub, mes, an), v in lus.items():
                cle_l = (typ, bloc, cat or None, rub, mes or None)
                pubs = d.setdefault(cle_l, {}).setdefault(an, [])
                pubs.append((lib, v, ref(f), f["ident"].get("Monnaie") or "F CFA"))
                npub = max(npub, len(pubs))
        donnees.append(d)

    e_id = ["No enregistrement", "No societe", "Edition", "Societe", "Pays", "Branche", "Monnaie", "Directeur general",
            "Date de creation", "Capital social (F CFA)", "Cadres", "Maitrise", "Employes", "Annee N", "Source"]
    idv = feuille("Identification", e_id)
    e_ep = ["No societe", "Societe", "Pays", "Branche", "Monnaie", "Bloc", "Categorie", "Rubrique", "Mesure", "Annee"]
    e_cc = ["No societe", "Societe", "Pays", "Branche", "Monnaie", "Rubrique", "Annee"]
    # Montants en milliers de F CFA, sauf les quelques fiches d'editions anciennes libellees dans une
    # autre monnaie (colonne Monnaie). Une annee est publiee dans une ou plusieurs editions (N dans l'une, N-1 dans la suivante) :
    # une paire de colonnes (valeur, edition) par publication, de la plus ancienne a la plus recente.
    # "Valeur" retient par formule la publication la plus recente qui donne un nombre ; a defaut,
    # "Non disponible" si la rubrique n'existe pas dans la mise en page des editions concernees.
    pub_b = [x for k in range(1, npub + 1) for x in (f"Valeur publication {k} (milliers)", f"Edition publication {k}")]
    pub_v = [x for k in range(1, npub + 1) for x in (f"Valeur publication {k}", f"Edition publication {k}")]
    epv = feuille("Emission&Prestations", e_ep + ["Valeur", "Source retenue"] + pub_v)
    ccv = feuille("Chiffres clés", e_cc + ["Valeur", "Source retenue"] + pub_v)
    epb = feuille("Emission&Prestations (brut)", e_ep + ["Valeur (milliers)", "Source retenue"] + pub_b
                  + [f"Reference publication {k}" for k in range(1, npub + 1)], True)
    ccb = feuille("Chiffres clés (brut)", e_cc + ["Valeur (milliers)", "Source retenue"] + pub_b
                  + [f"Reference publication {k}" for k in range(1, npub + 1)], True)

    def formules(n_base, r):
        cv = [get_column_letter(n_base + 3 + 2 * k) for k in range(npub)]
        cs = [get_column_letter(n_base + 4 + 2 * k) for k in range(npub)]
        plage = ",".join(f"{c}{r}" for c in cv)
        f_val = f'IF(COUNTIF({cv[0]}{r}:{cv[-1]}{r},"{ND}")>0,"{ND}","")' if npub > 1 else \
            f'IF({cv[0]}{r}="{ND}","{ND}","")'
        f_src = f'IF(COUNTIF({cv[0]}{r}:{cv[-1]}{r},"{ND}")>0,{cs[-1]}{r},"")' if npub > 1 else \
            f'IF({cv[0]}{r}="{ND}",{cs[0]}{r},"")'
        for c, e in zip(cv, cs):
            f_val = f"IF(ISNUMBER({c}{r}),{c}{r},{f_val})"
            f_src = f"IF(ISNUMBER({c}{r}),{e}{r},{f_src})"
        return "=" + f_val, "=" + f_src, cv, cs

    n_enr = 0
    for no, (s, d) in enumerate(zip(societes, donnees), 1):
        base_soc = [no, s["nom"], s["pays"], s["branche"]]
        for lib in libelles:
            f = s["fiches"].get(lib)
            if f is None:
                continue
            n_enr += 1
            i = f["ident"]
            idv.append([n_enr, no, lib, s["nom"] if f["illisible"] else i["Societe"], s["pays"],
                        s["branche"], i.get("Monnaie") or "F CFA", i["DG"] or None, ex.date_ou_texte(i["Date"]), ex.capital(i["Capital"]),
                        i["Cadres"], i["Maitrise"], i["Employes"], f["annees"][1], ref(f)])
        for cle_l in sorted(d, key=lambda c: tuple(x or "" for x in c)):
            typ, bloc, cat, rub, mes = cle_l
            for an in sorted(d[cle_l]):
                pubs = sorted(d[cle_l][an], key=lambda p: rang[p[0]])
                mon = " / ".join(dict.fromkeys(p[3] for p in pubs))
                if typ == "EP":
                    wsb, wsv, nomb, cles = epb, epv, "Emission&Prestations (brut)", [mon, bloc, cat, rub, mes, an]
                    n_base = len(e_ep)
                else:
                    wsb, wsv, nomb, cles = ccb, ccv, "Chiffres clés (brut)", [mon, rub, an]
                    n_base = len(e_cc)
                r = wsb.max_row + 1
                f_val, f_src, cv, cs = formules(n_base, r)
                vals = [x for lib, v, _, _ in pubs for x in (v, lib)] + [None, None] * (npub - len(pubs))
                refs = [p[2] for p in pubs] + [None] * (npub - len(pubs))
                wsb.append(base_soc + cles + [f_val, f_src] + vals + refs)
                cr = get_column_letter(n_base + 1)
                ligne_v = base_soc + cles + [f"=IF(ISNUMBER('{nomb}'!{cr}{r}),'{nomb}'!{cr}{r}*1000,'{nomb}'!{cr}{r})",
                                             f"='{nomb}'!{get_column_letter(n_base + 2)}{r}"]
                for c, e in zip(cv, cs):
                    ligne_v += [f"=IF(ISNUMBER('{nomb}'!{c}{r}),'{nomb}'!{c}{r}*1000,IF('{nomb}'!{c}{r}=\"\",\"\",'{nomb}'!{c}{r}))",
                                f"=IF('{nomb}'!{e}{r}=\"\",\"\",'{nomb}'!{e}{r})"]
                wsv.append(ligne_v)

    # --- Verification : bilan des tests de recalcul et concordance entre publications ---
    vf = feuille("Verification", ["Controle", "Source", "Resultat", "Detail"])
    for lib, nb_fiches, tot, ind, som, tests, echecs in (bilans or []):
        tot_ = max(tot, 1)
        vf.append(["Fiches lues", lib, nb_fiches, ""])
        vf.append(["Valeurs non nulles", lib, tot, ""])
        vf.append(["Verifiees individuellement", lib, f"{100 * ind / tot_:.1f} %",
                   f"{ind} valeurs recalculees a partir de l'evolution ou du ratio sinistres/primes imprimes"])
        vf.append(["Verifiees par un total", lib, f"{100 * som / tot_:.1f} %",
                   f"{som} valeurs dont la somme redonne le total imprime"])
        vf.append(["Non verifiees", lib, f"{100 * (tot - ind - som) / tot_:.1f} %",
                   f"{tot - ind - som} valeurs sans controle possible (valeur isolee, une seule annee...)"])
        vf.append(["Tests de recalcul en echec", lib, f"{len(echecs)} / {tests}",
                   "incoherences de la source (ratio ou evolution imprime faux), detail ci-dessous"])
    for nomb, d_type in (("Emission&Prestations", "EP"), ("Chiffres cles", "CC")):
        deux = egal = 0
        for d in donnees:
            for cle_l, par_an in d.items():
                if cle_l[0] != d_type:
                    continue
                for pubs in par_an.values():
                    vs = [p[1] for p in pubs if isinstance(p[1], (int, float))]
                    if len(vs) >= 2:
                        deux += 1
                        egal += max(vs) - min(vs) <= max(1, 0.001 * max(abs(x) for x in vs))
        if deux:
            vf.append([f"Concordance entre publications ({nomb})", "annees publiees deux fois",
                       f"{100 * egal / deux:.1f} %", f"{egal} valeurs identiques sur {deux} "
                       "(les ecarts restants sont des revisions entre deux publications)"])
    for lib, nb_fiches, tot, ind, som, tests, echecs in (bilans or []):
        for src, soc, x in echecs:
            vf.append(["Echec de recalcul", lib, f"{src} {soc}", x])

    nt = feuille("Notes", ["Type", "Reference", "Detail"])
    nt.append(["Monnaie", "toutes editions", "montants en milliers de la monnaie indiquee (F CFA sauf "
               "quelques fiches des editions 2005 a 2016 : Burundi, Guinee, Madagascar, Rwanda, Mauritanie...)"])
    nt.append(["Total affaires directes", "editions 12e, 18e, 19e...", "annee publiee sans detail par branche : "
               "le total imprime est repris, rubrique Total affaires directes"])
    nt.append([ND, "toutes editions", "la rubrique ou la mesure n'existe pas dans la mise en page de l'edition "
               "(nomenclature ou bloc different) ; une case vide signifie que la rubrique existe mais que "
               "la fiche ne la renseigne pas"])
    for (lib, g), cles in sorted(par_groupe.items(), key=lambda x: (rang.get(x[0][0], 0), x[0][1])):
        nt.append(["Structure", f"{lib} / {g}", f"{len(cles)} rubriques et mesures publiees"])
    for n in notes:
        nt.append(list(n))

    for c in idv["I"][1:]:
        c.number_format = "dd/mm/yyyy"
    for c in idv["J"][1:]:
        c.number_format = "#,##0"
    for ws, n_base in ((epv, len(e_ep)), (ccv, len(e_cc)), (epb, len(e_ep)), (ccb, len(e_cc))):
        fmt = "#,##0.000" if ws.sheet_state == "hidden" else "#,##0"
        for k in [n_base + 1] + [n_base + 3 + 2 * j for j in range(npub)]:
            for c in ws[get_column_letter(k)][1:]:
                c.number_format = fmt
    for ws in wb.worksheets:
        for col in ws.iter_cols(max_row=300):
            larg = max((len(str(c.value)) for c in col
                        if c.value is not None and not str(c.value).startswith("=")), default=12)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, larg + 2), 60)
    wb.save(sortie)


def lire_source(chemin):
    """PDF d'annuaire, ou dossiers de formulaires Excel separes par ";"."""
    if ";" in chemin or os.path.isdir(chemin):
        return lx.lire_dossiers([c for c in chemin.split(";") if c]) + ([],)
    reass = []
    fs, ign = ex.lire_pdf(chemin, reass)
    return fs, ign, reass


def main(sortie, editions_pdf):
    """editions_pdf : [(libelle, chemin du PDF)] de la plus ancienne a la plus recente."""
    editions = []
    notes_lecture = []
    with Pool(min(len(editions_pdf), os.cpu_count() or 1)) as pool:
        lus = pool.map(lire_source, [chemin for _, chemin in editions_pdf])
    for (lib, chemin), (fs, ign, reass) in zip(editions_pdf, lus):
        if reass:
            notes_lecture.append(("Hors champ", lib, "fiches de reassureurs, pages " + ", ".join(map(str, reass))))
        normaliser(fs)
        editions.append((lib, fs))
        if ign:
            notes_lecture.append(("Non lues", lib, ", ".join(map(str, ign))))
        attendu = Counter(f["annees"] for f in fs).most_common(1)[0][0] if fs else None
        for f in fs:
            for c in f["controles"]:
                notes_lecture.append(("Controle", f"{lib} : {ref(f)} {f['ident']['Societe'][:60]}", c))
            if f["annees"] != attendu:
                notes_lecture.append(("Exercices", f"{lib} : {ref(f)} {f['ident']['Societe']}",
                                      f"fiche publiee avec les exercices {f['annees'][0]}-{f['annees'][1]}"))
    societes, notes = construire(editions)
    notes += notes_lecture
    libelles = [lib for lib, _ in editions_pdf]
    bilans = []
    for lib, fs in editions:
        tot, ind, som, tests, echecs = vx.bilan(fs)
        bilans.append((lib, len(fs), tot, ind, som, tests, echecs))
    ecrire(societes, notes, sortie, libelles, bilans, editions)
    for lib, fs in editions:
        print(f"{lib} : {len(fs)} fiches")
    print(f"{len(societes)} societes")
    return societes, notes


if __name__ == "__main__":
    main(sys.argv[1], [tuple(a.split("=", 1)) for a in sys.argv[2:]])
